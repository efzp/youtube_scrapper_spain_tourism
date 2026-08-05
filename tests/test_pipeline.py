from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from src.config import Settings
from src.main import run_pipeline
from src.youtube.client import YouTubeRequestError


class FakeYouTubeClient:
    def search(self, **parameters: object) -> dict[str, object]:
        return {"items": [{"id": {"videoId": "same-video"}}]}

    def videos(self, **parameters: object) -> dict[str, object]:
        return {
            "items": [
                {
                    "id": "same-video",
                    "snippet": {"publishedAt": "2025-01-01T00:00:00Z"},
                    "contentDetails": {"duration": "PT2M"},
                    "statistics": {"viewCount": "10"},
                }
            ]
        }


class PermanentErrorClient:
    def __init__(self) -> None:
        self.calls = 0

    def search(self, **parameters: object) -> dict[str, object]:
        self.calls += 1
        raise YouTubeRequestError(403, "quotaExceeded")


class MemoryWriter:
    def __init__(self) -> None:
        self.tables: dict[str, list[dict[str, Any]]] = {}
        self.documents: dict[str, dict[str, Any]] = {}

    def write_table(
        self, name: str, records: list[dict[str, Any]], columns: list[str]
    ) -> None:
        self.tables[name] = records

    def write_json(self, name: str, payload: dict[str, Any]) -> None:
        self.documents[name] = payload


def _catalog(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "place_id", "municipio", "provincia", "comunidad_autonoma",
            "consulta_es_turismo", "consulta_en_review", "max_resultados_consulta",
            "activo", "prioridad",
        ]
    )
    sheet.append(
        ["ES-001", "Sevilla", "Sevilla", "Andalucia", "q1", "q2", 25, 1, 1]
    )
    sheet.add_table(Table(displayName="tbl_lugares", ref="A1:I2"))
    workbook.save(path)


def test_pipeline_summary_and_relational_deduplication(tmp_path: Path) -> None:
    catalog = tmp_path / "catalog.xlsx"
    _catalog(catalog)
    writer = MemoryWriter()
    fixed_now = datetime(2026, 1, 1, tzinfo=timezone.utc)
    settings = Settings(
        api_key="dummy-for-injected-client",
        catalog_path=catalog,
        output_dir=tmp_path,
        max_places=1,
    )

    summary = run_pipeline(
        settings,
        client=FakeYouTubeClient(),  # type: ignore[arg-type]
        writer=writer,
        now_factory=lambda: fixed_now,
    )

    assert summary.status == "COMPLETED"
    assert summary.queries_executed == 2
    assert summary.search_list_calls_estimated == 2
    assert summary.search_list_calls == 2
    assert summary.raw_results == 2
    assert summary.unique_videos == 1
    assert summary.videos_list_calls == 1
    assert summary.total_quota_units_estimated == 3
    assert len(writer.tables["youtube_search_results"]) == 2
    assert len(writer.tables["youtube_video_metadata"]) == 1
    assert writer.documents["youtube_run_summary"]["status"] == "COMPLETED"


def test_pipeline_stops_remaining_queries_after_permanent_error(tmp_path: Path) -> None:
    catalog = tmp_path / "catalog.xlsx"
    _catalog(catalog)
    writer = MemoryWriter()
    client = PermanentErrorClient()
    settings = Settings("dummy", catalog, tmp_path, max_places=1)

    summary = run_pipeline(
        settings,
        client=client,  # type: ignore[arg-type]
        writer=writer,
    )

    assert client.calls == 1
    assert summary.status == "FAILED"
    assert summary.queries_executed == 0
    assert summary.error_messages
